"""Тесты читаемой выдачи полного отчёта (ТЗ-41).

Проверяется: разбор плоского текста отчёта в структуру с САМОПРОВЕРКОЙ (число
разобранных пунктов блока сходится с «показано N», непонятая строка — ошибка),
Markdown-мастер, DOCX (Calibri 11, гиперссылки), неизменность stdout при новых
CLI-флагах и шаг сборки docx в обёртке. Живые проверки идут на КОПИИ боевой
базы (``--db <копия>``), боевая ``data/tuber.db`` не трогается.
"""
from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from tuber import config
from tuber.analysis import readable, report
from tuber.core import db, schema
from tuber.tools import db_backup

#: Синтетический полный отчёт по правилам §2.1 ТЗ-41 (раздел 1 + «Виральные»).
SYNTHETIC = "\n".join([
    "tuber report — объединённая выдача (единая база: (ядро))",
    "окно: последние 10 дней (с 2026-09-08 07:45:34 UTC), "
    "сформировано 2026-09-18 07:45:34 UTC",
    "=" * 72,
    "",
    "1. YouTube — топ по просмотрам в сутки за окно.",
    "   Правило: окно 100 видео с замерами.",
    "   Grok Bot за $20: девять сценариев использования — канал Paul J Lipsky; "
    "просмотры/сутки 1 912 008, просмотры 5 789 413, лайки 9 116; "
    "https://www.youtube.com/watch?v=UyMJBUCyDIs",
    "   Второе видео — канал Test Channel; просмотры/сутки 1 000, "
    "просмотры 2 000, лайки 10; https://example.com/v2",
    "",
    "5. Виральные — охват против своей аудитории.",
    "   Метрика: охват = просмотры ÷ подписчики СВОЕГО канала.",
    "   Блок 1. YouTube — полезное.",
    "   Блок 1: рассмотрено 100, прошло полы 10, показано 3.",
    "   Завод TerraFab утроит выпуск чипов для ИИ — канал Rayan Miller; "
    "подписчиков 1 140, просмотров 499 494, в сутки 7 892, "
    "охват на подписчика 438.15; опубликовано 8 дн. назад; "
    "тема: чипы и железо, язык en; "
    "https://www.youtube.com/watch?v=PrYfoYKqR24",
    "   Второй пункт — канал Канал Два; подписчиков 2 000, просмотров 100 000, "
    "в сутки 5 000, охват на подписчика 50.00; опубликовано 3 дн. назад; "
    "тема: кодинг и разработка, язык en; https://example.com/a",
    "   Третий пункт — канал Канал Три; подписчиков 3 000, просмотров 90 000, "
    "в сутки 4 000, охват на подписчика 30.00; опубликовано 2 дн. назад; "
    "тема: наука и медицина, язык ru; https://example.com/b",
    "      3.1. YouTube: просмотров ≥ 50 000, по просмотрам убыв.",
    "   3.1: рассмотрено 100, прошло полы 5, показано 2.",
    "   Продукт один — канал А; просмотров 60 000, в сутки 3 000, "
    "охват на подписчика 1.20; опубликовано 1 дн. назад; "
    "тема: запуски и анонсы, язык en; https://example.com/c",
    "   Продукт два — канал Б; просмотров 55 000, в сутки 2 000, "
    "охват на подписчика 0.90; опубликовано 2 дн. назад; "
    "тема: стартапы и бизнес, язык en; https://example.com/d",
    "      3.2. X: лайков ≥ 500, по лайкам убыв.",
    "   3.2: рассмотрено 50, прошло полы 4, показано 2.",
    "   @sama: Posts about agents — лайки 5 000, ответы 921; "
    "опубликовано 1 дн. назад; тема: релизы моделей, язык en; "
    "https://x.com/sama/status/1",
    "   @elonmusk: Grok now has a voice — лайки 4 000, ответы 100; "
    "опубликовано 0 дн. назад; тема: релизы моделей, язык en; "
    "https://x.com/elonmusk/status/2",
    "",
    "=" * 72,
    "ПОДВАЛ: честные оговорки (что в данных неполно)",
    "   YouTube: viral_index пуст у 1 из 2 строк score — индекс не построен.",
])


# --------------------------------------------------------------------------- #
# Разбор (1–3)
# --------------------------------------------------------------------------- #
def test_tz41_parse_top_section():
    doc = readable.parse_report(SYNTHETIC)
    assert [s.number for s in doc.sections] == [1, 5]
    assert len(doc.sections) == 2

    viral = doc.sections[1]
    block1 = next(b for b in viral.blocks if b.title.startswith("Блок 1."))
    assert block1.shown == 3
    assert len(block1.items) == 3
    subs = [b.title for b in viral.blocks if b.title[:4] in ("3.1.", "3.2.")]
    assert subs == ["3.1. YouTube: просмотров ≥ 50 000, по просмотрам убыв.",
                    "3.2. X: лайков ≥ 500, по лайкам убыв."]

    # У пункта поля отделены от текста, ссылка выделена.
    item = block1.items[0]
    assert item.text == "Завод TerraFab утроит выпуск чипов для ИИ"
    assert item.fields[0] == "канал Rayan Miller"
    assert any(f.startswith("подписчиков ") for f in item.fields)
    assert any(f.startswith("тема:") for f in item.fields)
    assert item.link == "https://www.youtube.com/watch?v=PrYfoYKqR24"

    # Кейс `@sama: … — лайки …` (правый кусок-метрика переносится в поля).
    x_block = next(b for b in viral.blocks if b.title.startswith("3.2."))
    assert x_block.items[0].text == "@sama: Posts about agents"
    assert x_block.items[0].fields[0] == "лайки 5 000, ответы 921"

    # Footer разобран отдельным разделом.
    assert doc.footer.title.startswith("ПОДВАЛ:")
    assert len(doc.footer.blocks[0].items) == 1


def test_tz41_parse_mismatch_raises():
    more = SYNTHETIC.replace("показано 3.", "показано 4.")
    with pytest.raises(ValueError, match="показано 4"):
        readable.parse_report(more)
    fewer = SYNTHETIC.replace("показано 3.", "показано 2.")
    with pytest.raises(ValueError):
        readable.parse_report(fewer)


def test_tz41_unknown_line_raises():
    broken = SYNTHETIC.replace(
        "5. Виральные — охват против своей аудитории.",
        "5. Виральные — охват против своей аудитории.\nНЕПОНЯТНАЯ СТРОКА БЕЗ РАЗДЕЛА")
    with pytest.raises(ValueError, match="непонятая строка"):
        readable.parse_report(broken)


# --------------------------------------------------------------------------- #
# Markdown (4)
# --------------------------------------------------------------------------- #
def test_tz41_markdown_layout():
    md = readable.render_markdown(readable.parse_report(SYNTHETIC))
    assert md.startswith("# Tuber — полная сводка: 18.09.2026")
    assert "## 1. YouTube" in md
    assert "### Блок 1." in md
    assert "**1. Завод TerraFab" in md
    assert "- ссылка: https://www.youtube.com/watch?v=PrYfoYKqR24" in md
    assert "=====" not in md
    assert "18.09.2026" in md
    for line in md.splitlines():
        assert not line.startswith(" "), f"ведущий пробел: {line!r}"


# --------------------------------------------------------------------------- #
# DOCX (5)
# --------------------------------------------------------------------------- #
def test_tz41_docx_written(tmp_path):
    path = tmp_path / "report.docx"
    readable.render_docx(readable.parse_report(SYNTHETIC), path)
    assert path.is_file() and path.stat().st_size > 0

    doc = Document(str(path))
    assert len(doc.paragraphs) >= 40
    links = sum(1 for _ in doc.element.body.iter(qn("w:hyperlink")))
    assert links >= 1

    bold = [p.text for p in doc.paragraphs
            if any(r.bold for r in p.runs)]
    assert any(t.startswith("1. YouTube") for t in bold)
    assert any(t.startswith("Tuber — полная сводка") for t in bold)

    runs = [r for p in doc.paragraphs for r in p.runs]
    assert any(r.font.name == "Calibri" and r.font.size is not None
               and r.font.size.pt == 11 for r in runs)
    assert doc.styles["Normal"].font.name == "Calibri"
    assert doc.styles["Normal"].font.size.pt == 11
    assert doc.sections[0].left_margin.cm == pytest.approx(2.0, abs=0.01)


# --------------------------------------------------------------------------- #
# CLI-флаги не меняют текст (6)
# --------------------------------------------------------------------------- #
def _live_copy(tmp_path: Path) -> str:
    env = os.environ.get("TUBER_TEST_LIVE_DB")
    if env and Path(env).is_file():
        return env
    backups = config.ROOT / "data" / "backups"
    if backups.is_dir():
        candidates = sorted(backups.glob("tuber-*.db"),
                            key=lambda p: p.stat().st_mtime, reverse=True)
        if candidates:
            return str(candidates[0])
    live = config.db_path(None)
    if Path(live).is_file():
        info = db_backup.backup_database(live, backup_dir=str(tmp_path), log=False)
        return info["backup"]
    pytest.skip("нет боевой базы и её копии")


@pytest.fixture(scope="module")
def live_copy(tmp_path_factory) -> str:
    """Копия боевой базы с дозаполненным ``content.url`` (как в соседних тестах)."""
    src = _live_copy(tmp_path_factory.mktemp("live41"))
    dst = tmp_path_factory.mktemp("live41_filled") / "live.db"
    shutil.copy2(src, dst)
    conn = db.connect(str(dst))
    try:
        with db.write_tx(conn):
            schema.ensure_content_url(conn)
    finally:
        conn.close()
    return str(dst)


def test_tz41_cli_flags_keep_text_output(live_copy, tmp_path, monkeypatch, capsys):
    """Текст ``--save`` и сводка ``--compact`` с флагами и без — байт-в-байт.

    Время фиксируется (``report.datetime``): иначе два прогона различаются
    секундой в «сформировано» и в «лайки/час», и побайтовое сравнение теряет
    смысл. Проверяется именно влияние новых флагов.
    """
    from datetime import datetime as _dt, timezone as _tz

    # Память выдачи отключена: два прогона с --save иначе намеренно
    # отличались бы (второй исключал бы показанное первым). Здесь проверяется
    # влияние флагов --readable/--docx на stdout, а не память (ТЗ-Tuber §2).
    monkeypatch.setenv("TUBER_DIGEST_MEMORY", "0")

    class _FrozenDatetime(_dt):
        @classmethod
        def now(cls, tz=None):
            return _dt(2026, 9, 18, 7, 45, 34, tzinfo=_tz.utc)

    monkeypatch.setattr(report, "datetime", _FrozenDatetime)

    out_txt = tmp_path / "full.txt"
    base = ["--db", live_copy, "--save", "--out", str(out_txt)]
    assert report.main(base) == 0
    text_plain = capsys.readouterr().out

    md = tmp_path / "full.md"
    docx = tmp_path / "full.docx"
    assert report.main(base + ["--readable", str(md), "--docx", str(docx)]) == 0
    text_flags = capsys.readouterr().out
    assert text_plain == text_flags
    assert md.is_file() and docx.is_file()

    assert report.main(["--db", live_copy, "--compact", "--save",
                        "--out", str(out_txt)]) == 0
    compact_plain = capsys.readouterr().out
    assert report.main(["--db", live_copy, "--compact", "--save",
                        "--out", str(out_txt), "--readable", str(md),
                        "--docx", str(docx)]) == 0
    compact_flags = capsys.readouterr().out
    assert compact_plain == compact_flags
    assert "сводка владельцу" in compact_plain


# --------------------------------------------------------------------------- #
# Обёртка (7)
# --------------------------------------------------------------------------- #
def test_tz41_wrapper_builds_docx():
    script = (config.ROOT / "scripts" / "common" / "tuber_report.sh").read_text(
        encoding="utf-8")
    assert "--docx" in script
    assert "--readable" in script
    assert ">/dev/null" in script or '2>>"$LOG"' in script
    # Прежние инварианты доставки сохранены.
    assert "--compact" in script
    assert "--save" in script
    assert "ALERT Tuber report" in script


# --------------------------------------------------------------------------- #
# Живой отчёт разбирается без потерь (8)
# --------------------------------------------------------------------------- #
def test_tz41_live_parse_no_loss(live_copy):
    conn = db.connect(live_copy, readonly=True)
    try:
        text = report.build(conn, days=10, db_path=live_copy)
    finally:
        conn.close()
    doc = readable.parse_report(text)
    shown = sum(b.shown for s in doc.sections for b in s.blocks
                if b.shown is not None)
    parsed = sum(len(b.items) for s in doc.sections for b in s.blocks
                 if b.shown is not None)
    assert shown > 0
    assert parsed == shown

    md = readable.render_markdown(doc)
    assert "## 1. YouTube" in md
    assert "=====" not in md
