"""Читаемая выдача полного отчёта: Markdown-мастер и DOCX (ТЗ-41).

Полный отчёт :mod:`tuber.analysis.report` печатается плоским текстом (абзацы,
отступы пробелами, поля в одну строку через ``;``) — «невозможно читать
глазами». Этот модуль разбирает тот же текст обратно в структуру
(:func:`parse_report`) и отдаёт его читаемым документом:

* :func:`render_markdown` — Markdown-мастер: заголовки разделов и блоков, пункт
  отдельным абзацем с жирным заголовком, поля пункта списком, ссылка отдельной
  строкой;
* :func:`render_docx` — DOCX через ``python-docx`` (без новых зависимостей):
  Calibri 11, чёрный текст, настоящие кликабельные гиперссылки, поля страницы
  2 см.

Текст ``build()`` и ``build_compact()`` НЕ меняется: модуль только читает его.

Самопроверка разбора обязательна (ТЗ-41 §3.1): число разобранных пунктов блока
сверяется с ``показано N`` из строки статистики; расхождение и любая непонятая
строка — :class:`ValueError`, а не тихая потеря пункта. Модуль не ходит в сеть и
не пишет в боевую базу.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

# --------------------------------------------------------------------------- #
# Разбор текста отчёта
# --------------------------------------------------------------------------- #
#: Известные начала полей пункта (ТЗ-41 §3.1). Сегмент, начинающийся с одного из
#: них, забирается из строки в поля пункта, а не остаётся в тексте.
FIELD_PREFIXES = (
    "канал ", "подписчиков ", "просмотров", "просмотры", "в сутки", "лайки",
    "ответы", "опубликовано", "тема:", "язык", "кратность", "норма", "охват",
    "база канала", "значимость", "https://", "http://",
)
#: Начала «метрик» — правый кусок после `` — `` переносится в поля только если
#: начинается с метрики (:data:`FIELD_PREFIXES` без чистых ссылок). ``канал``
#: включён: строка ``… — канал X`` несёт поле канала (ТЗ-41 §3.3).
METRIC_PREFIXES = tuple(p for p in FIELD_PREFIXES if not p.startswith("http"))

_SEP = "=" * 72
_SECTION_RE = re.compile(r"^(\d+)\. (.+)$")
_BLOCK_RE = re.compile(r"^   Блок (\d+)\. (.+)$")
_SUBBLOCK_RE = re.compile(r"^      (\d+\.\d+)\. (.+)$")
_BLOCK_STATS_RE = re.compile(
    r"^   Блок (\d+): рассмотрено (\d+), прошло полы (\d+), показано (\d+)\.$")
_SUB_STATS_RE = re.compile(
    r"^   (\d+\.\d+): рассмотрено (\d+), прошло полы (\d+), показано (\d+)\.$")
_NOTE_RE = re.compile(r"^\s+(Правило|Метрика|Свежесть|Оговорка|Границы):")
_SHOWN_CHANNELS_RE = re.compile(r"^\s+показаны \d+ каналов из \d+")
_STUB_RE = re.compile(r"^\s+нет ")
_URL_RE = re.compile(r"^https?://")


@dataclass
class Item:
    """Пункт раздела/блока: текст, поля и (необязательно) ссылка."""

    text: str
    fields: list[str] = field(default_factory=list)
    link: Optional[str] = None


@dataclass
class Block:
    """Блок раздела.

    ``title`` пуст у «неявного» блока (разделы 1–4 не печатают заголовок
    ``Блок N``); у блоков «Виральных» хранится как есть (``Блок 1. …`` либо
    ``3.1. …``). ``shown`` равен ``None``, когда строки статистики нет и
    сверять нечего. ``stats`` — исходная строка статистики, ``stub`` —
    заглушка «нет …» при ``показано 0``.
    """

    title: str = ""
    stats: str = ""
    shown: Optional[int] = None
    items: list[Item] = field(default_factory=list)
    stub: Optional[str] = None


@dataclass
class Section:
    """Раздел отчёта: заголовок, номер, пояснения и блоки."""

    title: str
    number: int
    notes: list[str] = field(default_factory=list)
    blocks: list[Block] = field(default_factory=list)


@dataclass
class ReportDoc:
    """Разобранный полный отчёт."""

    header: list[str]
    sections: list[Section]
    footer: Section


def _starts_with(seg: str, prefixes: tuple[str, ...]) -> bool:
    return any(seg.startswith(p) for p in prefixes)


def parse_item(raw: str) -> Item:
    """Разобрать строку пункта на текст, поля и ссылку (ТЗ-41 §3.1).

    Сегменты строки делятся по ``"; "``; с конца забираются те, что начинаются с
    известного поля. Остаток — текст пункта. Затем, если в тексте есть `` — `` и
    правый кусок начинается с метрики, он переносится в поля первым. Сегмент,
    начинающийся с ``http``, становится ссылкой.
    """
    body = raw.strip()
    segments = body.split("; ")
    idx = len(segments) - 1
    fields: list[str] = []
    while idx > 0 and _starts_with(segments[idx], FIELD_PREFIXES):
        fields.insert(0, segments[idx])
        idx -= 1
    text = "; ".join(segments[:idx + 1])

    # Кейс `@sama: … — лайки 3 831, ответы 921` (ТЗ-41 §3.1): хвост-метрика без
    # разделителя `; ` отделяется от текста. Берём ПОСЛЕДНЕЕ ` — `: описание
    # пункта само может содержать тире.
    if " — " in text:
        left, _, right = text.rpartition(" — ")
        if _starts_with(right, METRIC_PREFIXES):
            fields.insert(0, right)
            text = left

    link: Optional[str] = None
    kept: list[str] = []
    for f in fields:
        if _URL_RE.match(f):
            link = f
        else:
            kept.append(f)
    return Item(text=text.strip(), fields=kept, link=link)


class _Builder:
    """Пошаговый сборщик :class:`ReportDoc` с самопроверкой блоков."""

    def __init__(self) -> None:
        self.sections: list[Section] = []
        self.footer: Optional[Section] = None
        self.section: Optional[Section] = None
        self.block: Optional[Block] = None
        self.mode = ""  # "", "sections", "footer"

    # -- финализация ------------------------------------------------------- #
    def _close_block(self) -> None:
        blk = self.block
        if blk is None:
            return
        if blk.shown is not None and len(blk.items) != blk.shown:
            raise ValueError(
                f"разбор отчёта: блок «{blk.title or '(без названия)'}»: "
                f"статистика «показано {blk.shown}», а разобрано пунктов "
                f"{len(blk.items)} — расхождение (тихая потеря пункта).")
        self.block = None

    # -- переходы ---------------------------------------------------------- #
    def start_section(self, number: int, title: str) -> None:
        self._close_block()
        self.section = Section(title=title, number=number)
        self.sections.append(self.section)
        self.mode = "sections"

    def start_block(self, title: str) -> None:
        assert self.section is not None
        self._close_block()
        self.block = Block(title=title)
        self.section.blocks.append(self.block)

    def start_footer(self, title: str) -> None:
        self._close_block()
        self.footer = Section(title=title, number=0)
        self.section = self.footer
        self.block = Block()
        self.footer.blocks.append(self.block)
        self.mode = "footer"

    # -- наполнение -------------------------------------------------------- #
    def add_note(self, text: str) -> None:
        if self.section is None:
            raise ValueError(f"разбор отчёта: пояснение вне раздела: {text!r}")
        self.section.notes.append(text.strip())

    def add_stub(self, text: str) -> None:
        if self.block is not None and self.mode == "sections":
            self.block.stub = text.strip()
        elif self.section is not None:
            self.section.notes.append(text.strip())
        else:
            raise ValueError(f"разбор отчёта: заглушка вне раздела: {text!r}")

    def set_stats(self, raw: str, shown: int) -> None:
        if self.block is None:
            raise ValueError(f"разбор отчёта: статистика без блока: {raw!r}")
        self.block.stats = raw.strip()
        self.block.shown = shown

    def add_item(self, item: Item) -> None:
        if self.section is None:
            raise ValueError(f"разбор отчёта: пункт вне раздела: {item.text!r}")
        if self.block is None:
            self.block = Block()
            self.section.blocks.append(self.block)
        self.block.items.append(item)

    def extend_last_item(self, text: str) -> None:
        if self.block is None or not self.block.items:
            raise ValueError(f"разбор отчёта: продолжение без пункта: {text!r}")
        self.block.items[-1].fields.append(text.strip())


def parse_report(text: str) -> ReportDoc:
    """Разобрать текст полного отчёта в :class:`ReportDoc` (ТЗ-41 §3.1).

    Любая строка, не подошедшая ни под один вид, — :class:`ValueError` (лучше
    падение, чем тихая потеря пункта).
    """
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    if len(lines) < 3 or lines[2] != _SEP:
        raise ValueError("разбор отчёта: нет шапки (три строки, третья — '='*72)")
    header = lines[:3]

    b = _Builder()
    for raw in lines[3:]:
        if raw.strip() == "" or raw == _SEP:
            continue
        if raw.startswith("ПОДВАЛ:"):
            b.start_footer(raw.strip())
            continue
        m = _SECTION_RE.match(raw)
        if m:
            b.start_section(int(m.group(1)), m.group(2))
            continue
        if _BLOCK_RE.match(raw):
            b.start_block(raw.strip())
            continue
        if _SUBBLOCK_RE.match(raw):
            b.start_block(raw.strip())
            continue
        m = _BLOCK_STATS_RE.match(raw) or _SUB_STATS_RE.match(raw)
        if m:
            b.set_stats(raw, int(m.group(4)))
            continue
        if _NOTE_RE.match(raw) or _SHOWN_CHANNELS_RE.match(raw):
            b.add_note(raw)
            continue
        if _STUB_RE.match(raw):
            b.add_stub(raw)
            continue
        stripped = raw.strip()
        if raw.startswith("   ") and not raw.startswith("    "):
            b.add_item(parse_item(raw))
            continue
        if raw.startswith("     ") and b.block is not None:
            b.extend_last_item(stripped)
            continue
        raise ValueError(f"разбор отчёта: непонятая строка: {raw!r}")

    b._close_block()
    footer = b.footer or Section(title="", number=0)
    return ReportDoc(header=header, sections=b.sections, footer=footer)


# --------------------------------------------------------------------------- #
# Общие преобразования заголовка/полей
# --------------------------------------------------------------------------- #
def _header_lines(doc: ReportDoc) -> tuple[str, str]:
    """Заголовок документа и строка шапки из первых строк отчёта."""
    h0 = doc.header[0] if doc.header else ""
    h1 = doc.header[1] if len(doc.header) > 1 else ""
    date = ""
    m = re.search(r"сформировано (\d{4})-(\d{2})-(\d{2}) (\d{2}:\d{2}:\d{2}) UTC", h1)
    if m:
        date = f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    title = f"Tuber — полная сводка: {date}" if date else "Tuber — полная сводка"
    window = re.search(r"последние (\d+) дней", h1)
    cutoff = re.search(r"\(с ([^)]+?) UTC\)", h1)
    dbase = "(ядро)"
    if "(единая база: " in h0:
        dbase = h0.split("(единая база: ", 1)[1].rsplit(")", 1)[0]
    bits = []
    if window and cutoff:
        bits.append(f"Окно: последние {window.group(1)} дней (с {cutoff.group(1)} UTC).")
    if m:
        bits.append(f"Сформировано: {date} {m.group(4)} UTC.")
    bits.append(f"База: {dbase}.")
    return title, " ".join(bits)


def _field_text(field_: str) -> str:
    """Поле пункта для печати: ``канал X`` → ``канал: X`` (ТЗ-41 §3.3)."""
    if field_.startswith("канал "):
        return "канал: " + field_[len("канал "):]
    return field_


def _stats_text(block: Block) -> str:
    """Строка статистики блока без метки: ``Блок 1: рассмотрено …`` → ``Рассмотрено …``."""
    text = block.stats
    if ": " in text:
        text = text.split(": ", 1)[1]
    return text[:1].upper() + text[1:] if text else text


# --------------------------------------------------------------------------- #
# Markdown-мастер
# --------------------------------------------------------------------------- #
def render_markdown(doc: ReportDoc) -> str:
    """Собрать Markdown-мастер по макету ТЗ-41 §3.3."""
    title, subtitle = _header_lines(doc)
    out: list[str] = [f"# {title}", "", subtitle]
    for section in doc.sections:
        out.append("")
        out.append(f"## {section.number}. {section.title}")
        for note in section.notes:
            out.append("")
            out.append(f"*{note}*")
        for block in section.blocks:
            _md_block(out, block)
    if doc.footer and (doc.footer.title or doc.footer.blocks):
        out.append("")
        if doc.footer.title:
            out.append(f"## {doc.footer.title}")
        for note in doc.footer.notes:
            out.append("")
            out.append(f"*{note}*")
        for block in doc.footer.blocks:
            for item in block.items:
                out.extend(("", f"- {item.text}"))
    return "\n".join(out) + "\n"


def _md_block(out: list[str], block: Block) -> None:
    if block.title:
        out.append("")
        if re.match(r"^\d+\.\d+\. ", block.title):
            out.append(f"#### {block.title}")
        else:
            out.append(f"### {block.title}")
    if block.stats:
        out.append("")
        out.append(_stats_text(block))
    if block.stub:
        out.append("")
        out.append(block.stub)
    for i, item in enumerate(block.items, 1):
        out.append("")
        if block.title:
            out.append(f"**{i}. {item.text}**")
        else:
            out.append(f"### {i}. {item.text}")
        for f in item.fields:
            out.append(f"- {_field_text(f)}")
        if item.link:
            out.append(f"- ссылка: {item.link}")


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def _add_hyperlink(paragraph, url: str, text: str, size_pt: float = 11.0) -> None:
    """Добавить настоящую кликабельную гиперссылку в абзац.

    ``python-docx`` 1.2.0 не даёт ``add_hyperlink``, поэтому ``w:hyperlink``
    собирается вручную (ТЗ-41 §3.1, §3.4). Текст ссылки — полный URL, шрифт
    Calibri обычного размера, цвет чёрный (без «синего по умолчанию»).
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(size_pt * 2)))
    rpr.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _style_run(run, *, size: float, bold: bool = False, italic: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    run.font.size = _pt(size)
    run.font.color.rgb = _rgb_black()


def _pt(size: float):
    from docx.shared import Pt
    return Pt(size)


def _rgb_black():
    from docx.shared import RGBColor
    return RGBColor(0x00, 0x00, 0x00)


def render_docx(doc: ReportDoc, path) -> None:
    """Собрать DOCX по оформлению ТЗ-41 §3.4.

    Скучно и читаемо: Calibri 11, чёрный текст, никаких цветных заголовков,
    колонок и таблиц; настоящие гиперссылки; поля страницы 2 см.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    document = Document()
    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = _pt(11)
    normal.font.color.rgb = _rgb_black()

    title, subtitle = _header_lines(doc)
    p = document.add_paragraph()
    _style_run(p.add_run(title), size=16, bold=True)
    p = document.add_paragraph()
    _style_run(p.add_run(subtitle), size=10)

    for section in doc.sections:
        p = document.add_paragraph()
        _style_run(p.add_run(f"{section.number}. {section.title}"), size=14, bold=True)
        p.paragraph_format.page_break_before = True
        for note in section.notes:
            p = document.add_paragraph()
            _style_run(p.add_run(note), size=10, italic=True)
        for block in section.blocks:
            _docx_block(document, block)

    if doc.footer and (doc.footer.title or doc.footer.blocks):
        if doc.footer.title:
            p = document.add_paragraph()
            _style_run(p.add_run(doc.footer.title), size=14, bold=True)
            p.paragraph_format.page_break_before = True
        for block in doc.footer.blocks:
            for item in block.items:
                _docx_bullet(document, item.text)
    document.save(str(path))


def _docx_block(document, block: Block) -> None:
    from docx.shared import Cm

    if block.title:
        size = 11 if re.match(r"^\d+\.\d+\. ", block.title) else 12
        p = document.add_paragraph()
        _style_run(p.add_run(block.title), size=size, bold=True)
    if block.stats:
        p = document.add_paragraph()
        _style_run(p.add_run(_stats_text(block)), size=11)
    if block.stub:
        p = document.add_paragraph()
        _style_run(p.add_run(block.stub), size=11, italic=True)
    for i, item in enumerate(block.items, 1):
        p = document.add_paragraph()
        p.paragraph_format.space_after = _pt(6)
        _style_run(p.add_run(f"{i}. "), size=11, bold=True)
        _style_run(p.add_run(item.text), size=11)
        for f in item.fields:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            _style_run(p.add_run(f"• {_field_text(f)}"), size=11)
        if item.link:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            _style_run(p.add_run("• ссылка: "), size=11)
            _add_hyperlink(p, item.link, item.link, size_pt=11)
        document.add_paragraph()


def _docx_bullet(document, text: str) -> None:
    from docx.shared import Cm

    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    _style_run(p.add_run(f"• {text}"), size=11)


def write_readable(text: str, md_path: str | Path | None = None,
                   docx_path: str | Path | None = None) -> tuple[int, list[str]]:
    """Разобрать текст отчёта и записать мастер и/или DOCX.

    Возвращает ``(число пунктов, сообщения)``; сообщения — служебные строки для
    stderr (ТЗ-41 §3.2, §3.5). Ошибка разбора пробрасывается как ``ValueError``.
    """
    doc = parse_report(text)
    items = sum(len(b.items) for s in doc.sections for b in s.blocks)
    items += sum(len(b.items) for b in doc.footer.blocks) if doc.footer else 0
    messages: list[str] = []
    if md_path is not None:
        master = render_markdown(doc)
        p = Path(md_path)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(master, encoding="utf-8")
        messages.append(f"readable: {p}")
    if docx_path is not None:
        p = Path(docx_path)
        p.parent.mkdir(parents=True, exist_ok=True)
        render_docx(doc, p)
        messages.append(f"docx: {p}, {items} пунктов")
    return items, messages
