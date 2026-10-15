#!/usr/bin/env python3
"""Конвертер выкладки: markdown -> DOCX с настоящими гиперссылками (Calibri 11, чёрный текст)."""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1] if len(sys.argv) > 1 else "/root/tuber/reports/vykladka-2026-09-19.md"
DST = sys.argv[2] if len(sys.argv) > 2 else "/root/tuber/reports/vykladka-2026-09-19.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0, 0, 0)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
pf = style.paragraph_format
pf.space_after = Pt(4)
pf.line_spacing = 1.15

TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")
# Голый URL: всё до пробела/кавычки/угловой скобки.
BARE_URL = re.compile(r"https?://[^\s<>\"]+")
# Хвостовая пунктуация не входит в ссылку (кроме `)`, закрывающей markdown-ссылку,
# но такие URL сюда не попадают: их вырезает TOKEN в ветке `[...](...)`).
URL_TAIL = "),.;:!?\"'\u00bb"


def add_hyperlink(par, url, text):
    part = par.part
    r_id = part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(fonts)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); rPr.append(sz)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    hl.append(run)
    par._p.append(hl)


def add_plain(par, text):
    """Добавить обычный текст, превращая голые URL в настоящие гиперссылки."""
    pos = 0
    for m in BARE_URL.finditer(text):
        start, end = m.span()
        token = m.group(0)
        # URL внутри markdown-ссылки (конструкция `](...)`) не линкифицируем повторно.
        if text[max(0, start - 2):start] == "](":
            continue
        url = token.rstrip(URL_TAIL)
        if not url:
            continue
        if start > pos:
            par.add_run(text[pos:start])
        add_hyperlink(par, url, url)
        tail = token[len(url):]
        if tail:
            par.add_run(tail)
        pos = end
    if pos < len(text):
        par.add_run(text[pos:])


def add_runs(par, text):
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        elif piece.startswith("[") and "](" in piece:
            label, url = piece[1:-1].split("](", 1)
            add_hyperlink(par, url, label)
        else:
            add_plain(par, piece)
    for r in par.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)


def heading(text, level):
    sizes = {1: 17, 2: 14, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level > 1 else 0)
    p.paragraph_format.space_after = Pt(6)
    add_runs(p, text)
    for r in p.runs:
        r.bold = True
        r.font.size = Pt(sizes[level])
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p


lines = open(SRC, encoding="utf-8").read().splitlines()
for raw in lines:
    line = raw.rstrip()
    if not line.strip():
        continue
    if line.startswith("---"):
        continue
    if line.startswith("### "):
        heading(line[4:], 3)
    elif line.startswith("## "):
        heading(line[3:], 2)
    elif line.startswith("# "):
        heading(line[2:], 1)
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, line[2:])
    else:
        p = doc.add_paragraph()
        add_runs(p, line)

doc.save(DST)
print("saved:", DST)
print("paragraphs:", len(doc.paragraphs))
